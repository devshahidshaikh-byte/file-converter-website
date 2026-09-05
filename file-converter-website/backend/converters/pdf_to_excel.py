from pathlib import Path

import pdfplumber
from openpyxl import Workbook


def convert_pdf_to_excel(pdf_path: str, output_path: str) -> str:
    """
    Extract tables from a PDF and save them as an Excel workbook.
    """

    pdf_path = Path(pdf_path)
    output_path = Path(output_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    # Remove the default worksheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    with pdfplumber.open(str(pdf_path)) as pdf:

        for page_number, page in enumerate(pdf.pages, start=1):

            tables = page.extract_tables()

            if not tables:
                continue

            for table_number, table in enumerate(tables, start=1):

                # Create a worksheet for each table
                sheet_name = f"Page {page_number} Table {table_number}"

                # Excel sheet names cannot exceed 31 characters
                sheet_name = sheet_name[:31]

                worksheet = workbook.create_sheet(sheet_name)

                for row_number, row in enumerate(table, start=1):

                    for column_number, value in enumerate(row, start=1):

                        if value is None:
                            value = ""

                        worksheet.cell(
                            row=row_number,
                            column=column_number,
                            value=value
                        )

                # Make columns wider
                for column_cells in worksheet.columns:

                    max_length = 0
                    column_letter = column_cells[0].column_letter

                    for cell in column_cells:
                        if cell.value is not None:
                            max_length = max(
                                max_length,
                                len(str(cell.value))
                            )

                    worksheet.column_dimensions[
                        column_letter
                    ].width = min(max_length + 2, 40)

    # If no tables were found, create an informative worksheet
    if not workbook.worksheets:
        worksheet = workbook.create_sheet("No Tables Found")

        worksheet["A1"] = (
            "No tables were detected in this PDF. "
            "Please use a PDF containing selectable tables."
        )

    workbook.save(str(output_path))

    if not output_path.exists():
        raise RuntimeError("Excel file was not created.")

    return str(output_path)